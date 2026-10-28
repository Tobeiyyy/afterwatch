import base64
import io
import json
import sys
import urllib.error

import pytest

from afterwatch import config, exporters, writer

PNG = base64.b64decode(
    "iVBORw0KGgoAAAANSUhEUgAAAAEAAAABCAYAAAAfFcSJAAAADUlEQVR42mNkYPhfDwAChwGA60e6kgAAAABJRU5ErkJggg==")
VID = "dQw4w9WgXcQ"
NOTE = """# Some Title

**Tags:** [[youtube]] [[ai-tools|AI tools]]

## Key points
- First **bold** and *italic* and `code`
  - nested [site](https://example.com/a)
- Second

1. one
2. two

> quoted line

| Name | Value |
| --- | --- |
| a | 1 |
| b | 2 |

```python
print("hi")
```

![[shot-dQw4w9WgXcQ-0102.jpg]]
"""


def make_cfg(tmp_path, target, **output):
    cfg = config.load()
    cfg["output"].update({"target": target, "path": str(tmp_path / "out"), **output})
    return cfg


# --- inline / parse ---

def test_inline_styles_links_and_wikilinks():
    segs = exporters.inline("a **b** *c* `d` [e](https://x.io) [[t|alias]] [[plain]]")
    styled = [(s.text, s.bold, s.italic, s.code, s.url) for s in segs if s.bold or s.italic or s.code or s.url]
    assert styled == [("b", True, False, False, ""), ("c", False, True, False, ""),
                      ("d", False, False, True, ""), ("e", False, False, False, "https://x.io")]
    # wikilinks collapse into plain text
    assert "".join(s.text for s in segs) == "a b c d e alias plain"


def test_inline_leaves_lone_asterisks_and_empty_text():
    assert [s.text for s in exporters.inline("5 * 3 * 2")] == ["5 * 3 * 2"]
    assert exporters.inline("") == []


def test_inline_nests_bold_around_link():
    seg = exporters.inline("**[x](https://y.z)**")[0]
    assert seg.bold and seg.url == "https://y.z"


def test_parse_block_kinds():
    kinds = [(b.kind, b.level, b.num) for b in exporters.parse(NOTE)]
    assert kinds == [
        ("heading", 1, 0), ("para", 0, 0), ("heading", 2, 0),
        ("bullet", 0, 0), ("bullet", 1, 0), ("bullet", 0, 0),
        ("number", 0, 1), ("number", 0, 2), ("quote", 0, 0), ("table", 0, 0),
        ("code", 0, 0), ("image", 0, 0)]
    table = exporters.parse(NOTE)[9]
    assert table.rows == (("Name", "Value"), ("a", "1"), ("b", "2"))
    assert exporters.parse(NOTE)[-1].text == "shot-dQw4w9WgXcQ-0102.jpg"


def test_parse_numbering_restarts_after_other_block():
    blocks = exporters.parse("1. a\n2. b\n\ntext\n\n1. c")
    assert [b.num for b in blocks if b.kind == "number"] == [1, 2, 1]


def test_parse_boundaries():
    assert exporters.parse("") == []
    code = exporters.parse("```\nnever closed\nmore")[0]  # unclosed fence runs to the end
    assert (code.kind, code.text) == ("code", "never closed\nmore")
    assert exporters.parse("| a | b |\nplain")[0].kind == "para"  # no separator row: not a table
    assert exporters.parse("---")[0].kind == "hr"


# --- get() ---

def test_get_picks_exporter_by_target(tmp_path):
    for target, cls, hubs in [("obsidian", exporters.ObsidianExporter, True),
                              ("markdown", exporters.MarkdownExporter, True),
                              ("docx", exporters.DocxExporter, False),
                              ("txt", exporters.TxtExporter, False)]:
        exp = exporters.get(make_cfg(tmp_path, target))
        assert type(exp) is cls and exp.name == target and exp.supports_hubs is hubs
        assert exp.root == tmp_path / "out"


def test_get_rejects_unknown_target_and_missing_path(tmp_path):
    with pytest.raises(config.ConfigError, match="unknown output target 'pdf'"):
        exporters.get(make_cfg(tmp_path, "pdf"))
    with pytest.raises(config.ConfigError, match="output.path"):
        exporters.get(make_cfg(tmp_path, "txt", path="  "))


def test_get_defaults_to_obsidian_and_needs_a_path():
    with pytest.raises(config.ConfigError, match="output.path"):
        exporters.get()


# --- obsidian ---

def test_obsidian_is_byte_identical_to_writer(tmp_path):
    exp = exporters.get(make_cfg(tmp_path, "obsidian"))
    location = exp.write_note("Some Title", VID, NOTE)
    expected = writer.write_note(tmp_path / "ref", "Some Title", VID, NOTE)
    assert open(location, "rb").read() == expected.read_bytes()
    assert location.endswith("Some Title.md")


def test_obsidian_find_existing_tags_and_failures(tmp_path):
    exp = exporters.get(make_cfg(tmp_path, "obsidian"))
    assert exp.find_existing(VID) is None
    location = exp.write_note("Some Title", VID, NOTE)
    assert exp.find_existing(VID) == location
    assert exp.write_tags(["a", "b"]) == ["a", "b"]
    assert exp.write_tags(["a"]) == []  # already a stub
    exp.record_failure("https://youtu.be/x", "no captions")
    failures = (tmp_path / "out" / config.folder("failures")).read_text(encoding="utf-8")
    assert "https://youtu.be/x" in failures and "no captions" in failures


# --- markdown ---

def test_markdown_rewrites_images_and_references(tmp_path):
    exp = exporters.get(make_cfg(tmp_path, "markdown"))
    path = exp.write_note("Some Title", VID, NOTE)
    text = open(path, encoding="utf-8").read()
    assert "![](../Images/shot-dQw4w9WgXcQ-0102.jpg)" in text
    assert "![[" not in text and "<iframe" not in text
    assert text.endswith(f"# References\n[Watch on YouTube](https://www.youtube.com/watch?v={VID})\n")
    assert "[[ai-tools|AI tools]]" in text  # wikilinks stay
    assert exp.find_existing(VID) == path


def test_markdown_encodes_spaces_and_keeps_notes_without_images(tmp_path):
    cfg = make_cfg(tmp_path, "markdown")
    cfg["folders"]["images"] = "My Images"
    exp = exporters.get(cfg)
    text = open(exp.write_note("T", VID, "![[a b.jpg|300]]"), encoding="utf-8").read()
    assert "![](../My%20Images/a%20b.jpg)" in text
    plain = open(exp.write_note("T2", "BBBBBBBBBBB", "just text"), encoding="utf-8").read()
    assert plain.startswith("just text\n\n# References")


# --- txt ---

def test_txt_renders_plain_text(tmp_path):
    exp = exporters.get(make_cfg(tmp_path, "txt"))
    text = open(exp.write_note("Some Title", VID, NOTE), encoding="utf-8").read()
    assert text.startswith("Some Title\n==========\n")
    assert "Key points\n----------" in text
    assert "- First bold and italic and code" in text
    assert "  - nested site (https://example.com/a)" in text
    assert "1. one\n2. two" in text
    assert "> quoted line" in text
    assert "Name | Value\n-----+------\na    | 1" in text
    assert '    print("hi")' in text
    assert "[Screenshot: shot-dQw4w9WgXcQ-0102.jpg]" in text
    assert "[[" not in text and "**" not in text
    assert text.endswith(f"Watch on YouTube: https://www.youtube.com/watch?v={VID}\n")
    assert "Tags: youtube AI tools" in text


def test_txt_duplicates_and_title_collisions(tmp_path):
    exp = exporters.get(make_cfg(tmp_path, "txt"))
    assert exp.find_existing(VID) is None
    first = exp.write_note("Same", VID, "x")
    assert exp.find_existing(VID) == first
    second = exp.write_note("Same", "BBBBBBBBBBB", "y")
    assert second.endswith("Same (BBBBBBBBBBB).txt") and second != first
    (tmp_path / "out" / "Videos" / "Same.txt").unlink()
    assert exp.find_existing(VID) is None  # file gone: not a duplicate any more
    assert exp.find_existing("never-seen") is None


def test_txt_table_with_ragged_row_and_no_body(tmp_path):
    text = exporters.render_txt("T", "| a | b |\n|---|---|\n| only |", VID)
    assert "a    | b" in text and "only |" in text
    assert exporters.render_txt("T", "", VID).startswith("T\n=\n")


def test_txt_record_failure(tmp_path):
    exp = exporters.get(make_cfg(tmp_path, "txt"))
    exp.record_failure("u", "r")
    assert "u — r" in (tmp_path / "out" / config.folder("failures")).read_text(encoding="utf-8")


# --- docx ---

def _docx_note(tmp_path, note=NOTE, image=True):
    exp = exporters.get(make_cfg(tmp_path, "docx"))
    if image:
        images = tmp_path / "out" / "Images"
        images.mkdir(parents=True)
        (images / "shot-dQw4w9WgXcQ-0102.jpg").write_bytes(PNG)
    import docx
    return exp, docx.Document(exp.write_note("Some Title", VID, note))


def test_docx_structure(tmp_path):
    exp, doc = _docx_note(tmp_path)
    styled = [(p.style.name, p.text) for p in doc.paragraphs]
    assert ("Title", "Some Title") in styled
    assert ("Heading 2", "Key points") in styled
    assert ("List Bullet", "First bold and italic and code") in styled
    assert ("List Bullet 2", "nested site") in styled
    assert ("Quote", "quoted line") in styled
    assert any(t.startswith("1. one") for _, t in styled) and any(t.startswith("2. two") for _, t in styled)
    assert "print" in " ".join(t for _, t in styled)
    assert [[c.text for c in r.cells] for r in doc.tables[0].rows] == [["Name", "Value"], ["a", "1"], ["b", "2"]]
    assert len(doc.inline_shapes) == 1
    assert round(doc.inline_shapes[0].width.inches) == 6
    assert styled[-1][1] == f"Watch on YouTube: https://www.youtube.com/watch?v={VID}"
    assert exp.find_existing(VID).endswith("Some Title.docx")


def test_docx_formatting_and_hyperlinks(tmp_path):
    _, doc = _docx_note(tmp_path)
    runs = {r.text: r for p in doc.paragraphs for r in p.runs}
    assert runs["bold"].bold and runs["italic"].italic and runs["code"].font.name == "Consolas"
    targets = {rel.target_ref for rel in doc.part.rels.values() if rel.is_external}
    assert "https://example.com/a" in targets


def test_docx_missing_image_becomes_placeholder(tmp_path):
    _, doc = _docx_note(tmp_path, image=False)
    assert len(doc.inline_shapes) == 0
    placeholder = [r for p in doc.paragraphs for r in p.runs if "not found" in r.text][0]
    assert placeholder.italic and "shot-dQw4w9WgXcQ-0102.jpg" in placeholder.text


def test_docx_empty_note_and_unreadable_image(tmp_path):
    exp = exporters.get(make_cfg(tmp_path, "docx"))
    (tmp_path / "out" / "Images").mkdir(parents=True)
    (tmp_path / "out" / "Images" / "bad.jpg").write_bytes(b"not an image")
    import docx
    doc = docx.Document(exp.write_note("T", VID, "![[bad.jpg]]"))
    assert len(doc.inline_shapes) == 0 and any("not found" in p.text for p in doc.paragraphs)
    assert docx.Document(exp.write_note("E", "BBBBBBBBBBB", "")).paragraphs[0].text == "E"


def test_docx_without_python_docx_names_the_extra(tmp_path, monkeypatch):
    monkeypatch.setitem(sys.modules, "docx", None)
    with pytest.raises(config.ConfigError, match=r"pip install afterwatch\[docx\]"):
        exporters.get(make_cfg(tmp_path, "docx"))


# --- notion ---

PAGE_ID = "1429989fe8ac4effbc8f57f56486db54"
DASHED = "1429989f-e8ac-4eff-bc8f-57f56486db54"


class FakeResp(io.BytesIO):
    def __enter__(self):
        return self

    def __exit__(self, *exc):
        return False


class FakeNotion:
    """Stands in for urlopen: records requests, answers from a script."""

    def __init__(self, *script):
        self.script = list(script)
        self.requests = []

    def __call__(self, req, timeout=None):
        self.requests.append(req)
        item = self.script.pop(0) if self.script else None
        if isinstance(item, Exception):
            raise item
        return FakeResp(json.dumps(item or self.default(req)).encode())

    @staticmethod
    def default(req):
        if req.full_url.endswith("/file_uploads"):
            return {"id": "up-1", "status": "pending"}
        if req.full_url.endswith("/pages"):
            return {"id": "page-1", "url": "https://www.notion.so/page-1"}
        return {"status": "uploaded"}

    def bodies(self):
        return [json.loads(r.data) for r in self.requests if r.get_header("Content-type") == "application/json"]


def http_error(code, message="boom", retry_after=None):
    headers = {"Retry-After": retry_after} if retry_after else {}
    return urllib.error.HTTPError("https://api.notion.com", code, "x", headers,
                                  io.BytesIO(json.dumps({"message": message}).encode()))


@pytest.fixture
def notion(tmp_path, monkeypatch):
    monkeypatch.setenv("NOTION_TOKEN", "secret-token")
    sleeps = []
    monkeypatch.setattr(exporters.time, "sleep", sleeps.append)

    def build(*script, parent=PAGE_ID):
        fake = FakeNotion(*script)
        monkeypatch.setattr(exporters.urllib.request, "urlopen", fake)
        exp = exporters.get(make_cfg(tmp_path, "notion", path="", notion_parent=parent))
        return exp, fake, sleeps
    return build


def test_notion_parent_id_forms():
    for value in (PAGE_ID, DASHED, f"https://www.notion.so/ws/My-Page-{PAGE_ID}?pvs=4",
                  f"https://www.notion.so/{PAGE_ID}/", f"  {PAGE_ID.upper()}  "):
        assert exporters._notion_id(value) == DASHED
    for bad in ("", "not-an-id", PAGE_ID[:-1], "https://www.notion.so/ws/My-Page"):
        with pytest.raises(config.ConfigError, match="notion_parent"):
            exporters._notion_id(bad)


def test_notion_needs_token(tmp_path, monkeypatch):
    monkeypatch.delenv("NOTION_TOKEN", raising=False)
    monkeypatch.setattr(config, "api_key", lambda *a, **k: None)
    with pytest.raises(config.ConfigError, match="NOTION_TOKEN"):
        exporters.get(make_cfg(tmp_path, "notion", path="", notion_parent=PAGE_ID))


def test_notion_token_falls_back_to_config_store(tmp_path, monkeypatch):
    monkeypatch.delenv("NOTION_TOKEN", raising=False)
    monkeypatch.setattr(config, "api_key", lambda provider, cfg=None: "from-store")
    exp = exporters.get(make_cfg(tmp_path, "notion", path="", notion_parent=PAGE_ID))
    assert exp.token == "from-store" and exp.supports_hubs is False
    assert exp.root == config.state_dir() / "notion-images"


def test_notion_creates_page_with_headers_and_blocks(notion):
    exp, fake, _ = notion()
    url = exp.write_note("Some Title", VID, NOTE.replace("![[shot-dQw4w9WgXcQ-0102.jpg]]", ""))
    assert url == "https://www.notion.so/page-1"
    (req,) = fake.requests
    assert req.full_url == "https://api.notion.com/v1/pages" and req.get_method() == "POST"
    assert req.get_header("Authorization") == "Bearer secret-token"
    assert req.get_header("Notion-version") == "2026-03-11"
    body = json.loads(req.data)
    assert body["parent"] == {"page_id": DASHED}
    assert body["properties"]["title"]["title"][0]["text"]["content"] == "Some Title"
    blocks = body["children"]
    assert [b["type"] for b in blocks] == [
        "paragraph", "heading_2", "bulleted_list_item", "bulleted_list_item",
        "numbered_list_item", "numbered_list_item", "quote", "table", "code", "paragraph"]
    first = blocks[2]["bulleted_list_item"]
    assert [i["text"]["content"] for i in first["rich_text"]] == ["First ", "bold", " and ", "italic", " and ", "code"]
    assert first["rich_text"][1]["annotations"] == {"bold": True, "italic": False, "code": False}
    nested = first["children"][0]["bulleted_list_item"]["rich_text"]
    assert nested[1]["text"]["link"] == {"url": "https://example.com/a"}
    table = blocks[7]["table"]
    assert table["table_width"] == 2 and table["has_column_header"] is True
    assert [c[0]["text"]["content"] for c in table["children"][1]["table_row"]["cells"]] == ["a", "1"]
    assert blocks[8]["code"]["language"] == "python"
    last = blocks[-1]["paragraph"]["rich_text"][0]["text"]
    assert last["link"]["url"] == f"https://www.youtube.com/watch?v={VID}"
    assert exp.find_existing(VID) == "https://www.notion.so/page-1"


def test_notion_batches_blocks_by_100(notion):
    exp, fake, _ = notion()
    exp.write_note("Big", VID, "\n\n".join(f"p{i}" for i in range(250)))  # 250 + link paragraph
    assert [r.get_method() for r in fake.requests] == ["POST", "PATCH", "PATCH"]
    assert fake.requests[1].full_url == "https://api.notion.com/v1/blocks/page-1/children"
    assert [len(b.get("children", [])) for b in fake.bodies()] == [100, 100, 51]


def test_notion_batch_boundary_exactly_100(notion):
    exp, fake, _ = notion()
    exp.write_note("Edge", VID, "\n\n".join(f"p{i}" for i in range(99)))  # 99 + link = 100
    assert len(fake.requests) == 1 and len(fake.bodies()[0]["children"]) == 100


def test_notion_splits_long_text_at_2000(notion):
    exp, fake, _ = notion()
    exp.write_note("Long", VID, "x" * 4500 + "\n\n" + "y" * 2000)
    blocks = fake.bodies()[0]["children"]
    assert [len(i["text"]["content"]) for i in blocks[0]["paragraph"]["rich_text"]] == [2000, 2000, 500]
    assert [len(i["text"]["content"]) for i in blocks[1]["paragraph"]["rich_text"]] == [2000]  # exact fit


def test_notion_nesting_is_capped_at_one_level_and_unknown_code_language(notion):
    exp, fake, _ = notion()
    exp.write_note("N", VID, "- a\n  - b\n    - c\n\n```brainfuck\n+\n```")
    blocks = fake.bodies()[0]["children"]
    assert [c["bulleted_list_item"]["rich_text"][0]["text"]["content"]
            for c in blocks[0]["bulleted_list_item"]["children"]] == ["b", "c"]
    assert "children" not in blocks[0]["bulleted_list_item"]["children"][0]["bulleted_list_item"]
    assert blocks[1]["code"]["language"] == "plain text"


def test_notion_uploads_screenshots(notion):
    exp, fake, _ = notion()
    images = exp.root / "Images"
    images.mkdir(parents=True)
    (images / "shot-dQw4w9WgXcQ-0102.jpg").write_bytes(PNG)
    exp.write_note("Pic", VID, "![[shot-dQw4w9WgXcQ-0102.jpg]]")
    create, send, page = fake.requests
    assert create.full_url.endswith("/v1/file_uploads")
    assert json.loads(create.data) == {"filename": "shot-dQw4w9WgXcQ-0102.jpg", "content_type": "image/jpeg"}
    assert send.full_url.endswith("/v1/file_uploads/up-1/send")
    assert send.get_header("Content-type").startswith("multipart/form-data; boundary=")
    assert b'name="file"; filename="shot-dQw4w9WgXcQ-0102.jpg"' in send.data and PNG in send.data
    assert send.get_header("Notion-version") == "2026-03-11"
    assert json.loads(page.data)["children"][0] == {
        "object": "block", "type": "image",
        "image": {"type": "file_upload", "file_upload": {"id": "up-1"}}}


def test_notion_missing_screenshot_is_a_placeholder_without_upload(notion):
    exp, fake, _ = notion()
    exp.write_note("Pic", VID, "![[gone.jpg]]")
    assert len(fake.requests) == 1
    block = fake.bodies()[0]["children"][0]["paragraph"]["rich_text"][0]
    assert "gone.jpg" in block["text"]["content"] and block["annotations"]["italic"]


@pytest.mark.parametrize("code", [401, 403])
def test_notion_auth_errors_are_config_errors(notion, code):
    exp, _, _ = notion(http_error(code, "API token is invalid."))
    with pytest.raises(config.ConfigError, match=r"shared with your integration.*API token is invalid"):
        exp.write_note("T", VID, "x")
    assert exp.find_existing(VID) is None  # nothing recorded for a failed export


def test_notion_validation_error_carries_notions_message(notion):
    exp, _, _ = notion(http_error(400, "body.children[0].paragraph is not valid"))
    with pytest.raises(exporters.ExportError, match="400.*body.children"):
        exp.write_note("T", VID, "x")


def test_notion_retries_rate_limits_with_retry_after(notion):
    exp, fake, sleeps = notion(http_error(429, retry_after="2"), http_error(429, retry_after="500"))
    assert exp.write_note("T", VID, "x") == "https://www.notion.so/page-1"
    assert len(fake.requests) == 3 and sleeps == [2, 60]  # waits are capped at a minute


def test_notion_gives_up_after_a_few_rate_limits(notion):
    exp, fake, sleeps = notion(*[http_error(429, "slow down") for _ in range(4)])
    with pytest.raises(exporters.ExportError, match="429.*slow down"):
        exp.write_note("T", VID, "x")
    assert len(fake.requests) == 4 and len(sleeps) == 3


def test_notion_unreachable(notion):
    exp, _, _ = notion(urllib.error.URLError("no route"))
    with pytest.raises(exporters.ExportError, match="could not reach Notion"):
        exp.write_note("T", VID, "x")


def test_notion_failures_go_to_state_dir(notion):
    exp, _, _ = notion()
    exp.record_failure("u", "r")
    assert "u — r" in (config.state_dir() / config.folder("failures")).read_text(encoding="utf-8")
