"""Where finished notes go: an Obsidian vault, a plain Markdown folder, Word or
text files, or Notion. get() picks the exporter named in [output]."""
from __future__ import annotations

import json
import mimetypes
import os
import re
import time
import urllib.error
import urllib.parse
import urllib.request
import uuid
from dataclasses import dataclass, replace
from pathlib import Path

from . import config, writer

# Checked against developers.notion.com on 2026-10-28 (File Upload API needs this version).
NOTION_VERSION = "2026-03-11"
NOTION_API = "https://api.notion.com/v1"
_NOTION_TEXT_LIMIT = 2000   # characters per rich-text item
_NOTION_BATCH = 100         # top-level blocks per request


class ExportError(Exception):
    pass


# --- Markdown parsing shared by the docx, txt and Notion exporters ---

@dataclass
class Block:
    kind: str           # heading para bullet number code quote table image hr
    text: str = ""
    level: int = 0      # heading level, or list depth
    num: int = 0        # item number for "number"
    lang: str = ""
    rows: tuple = ()


@dataclass
class Seg:
    text: str
    bold: bool = False
    italic: bool = False
    code: bool = False
    url: str = ""


_FENCE = re.compile(r"^\s*(```|~~~)\s*([\w+#.-]*)")
_HEAD = re.compile(r"^(#{1,6})\s+(.*?)\s*$")
_ITEM = re.compile(r"^(\s*)([-*+]|\d+[.)])\s+(.*)$")
_HR = re.compile(r"^\s*([-*_])(\s*\1){2,}\s*$")
_SEP = re.compile(r"^\s*\|?\s*:?-+:?\s*(\|\s*:?-+:?\s*)*\|?\s*$")
_EMBED = re.compile(r"!\[\[([^\]|]+)(?:\|[^\]]*)?\]\]")
_WIKI = re.compile(r"\[\[([^\]|]+)(?:\|([^\]]*))?\]\]")
_INLINE = re.compile(r"`([^`]+)`|\*\*(.+?)\*\*|\*(?=\S)(.+?)(?<=\S)\*|\[([^\]]+)\]\(([^)\s]+)\)")


def _cells(line: str) -> list[str]:
    return [c.strip().replace("\\|", "|") for c in re.split(r"(?<!\\)\|", line.strip().strip("|"))]


def parse(md: str) -> list[Block]:
    """Split Markdown into blocks. Covers what the note prompt produces; anything
    unrecognised falls through as a paragraph."""
    lines = md.replace("\r\n", "\n").split("\n")
    out: list[Block] = []
    para: list[str] = []
    stack: list[int] = []   # indents of the open list levels
    counts: list[int] = []  # running number per list depth
    i = 0

    def flush():
        if para:
            out.append(Block("para", "\n".join(para)))
            para.clear()

    while i < len(lines):
        line = lines[i]
        i += 1
        if m := _FENCE.match(line):
            flush(); stack.clear(); counts.clear()
            body = []
            while i < len(lines) and not lines[i].strip().startswith(m[1]):
                body.append(lines[i])
                i += 1
            i += 1
            out.append(Block("code", "\n".join(body), lang=m[2].lower()))
        elif not line.strip():
            flush()
        elif m := _HEAD.match(line):
            flush(); stack.clear(); counts.clear()
            out.append(Block("heading", m[2], level=len(m[1])))
        elif _HR.match(line):
            flush(); stack.clear(); counts.clear()
            out.append(Block("hr"))
        elif (m := _EMBED.fullmatch(line.strip())):
            flush(); stack.clear(); counts.clear()
            out.append(Block("image", m[1].strip()))
        elif line.lstrip().startswith("|") and i < len(lines) and _SEP.match(lines[i]):
            flush(); stack.clear(); counts.clear()
            rows = [_cells(line)]
            i += 1
            while i < len(lines) and lines[i].lstrip().startswith("|"):
                rows.append(_cells(lines[i]))
                i += 1
            out.append(Block("table", rows=tuple(tuple(r) for r in rows)))
        elif m := _ITEM.match(line):
            flush()
            indent = len(m[1].expandtabs(4))
            while stack and indent < stack[-1]:
                stack.pop()
            if not stack or indent > stack[-1]:
                stack.append(indent)
            depth = len(stack) - 1
            del counts[depth + 1:]
            counts += [0] * (depth + 1 - len(counts))
            if m[2][0].isdigit():
                counts[depth] += 1
                out.append(Block("number", m[3], level=depth, num=counts[depth]))
            else:
                counts[depth] = 0
                out.append(Block("bullet", m[3], level=depth))
        elif line.lstrip().startswith(">"):
            flush(); stack.clear(); counts.clear()
            quote = [re.sub(r"^\s*>\s?", "", line)]
            while i < len(lines) and lines[i].lstrip().startswith(">"):
                quote.append(re.sub(r"^\s*>\s?", "", lines[i]))
                i += 1
            out.append(Block("quote", "\n".join(quote).strip()))
        elif stack and line[0] in " \t" and out and out[-1].kind in ("bullet", "number"):
            out[-1].text += "\n" + line.strip()  # wrapped list item
        else:
            stack.clear(); counts.clear()
            para.append(line.strip())
    flush()
    return out


def inline(text: str) -> list[Seg]:
    """Split one line of Markdown into styled runs. Wikilinks become plain text."""
    text = _EMBED.sub(lambda m: f"[Screenshot: {m[1].strip()}]", text)
    text = _WIKI.sub(lambda m: m[2] or m[1], text)
    return _inline(text, Seg(""))


def _inline(text: str, base: Seg) -> list[Seg]:
    segs: list[Seg] = []
    pos = 0
    for m in _INLINE.finditer(text):
        if m.start() > pos:
            segs.append(replace(base, text=text[pos:m.start()]))
        if m[1]:
            segs.append(replace(base, text=m[1], code=True))
        elif m[2]:
            segs += _inline(m[2], replace(base, bold=True))
        elif m[3]:
            segs += _inline(m[3], replace(base, italic=True))
        else:
            segs += _inline(m[4], replace(base, url=m[5]))
        pos = m.end()
    if pos < len(text):
        segs.append(replace(base, text=text[pos:]))
    return segs


def _plain(text: str) -> str:
    """Text without markup; links read as 'text (url)'."""
    return "".join(
        f"{s.text} ({s.url})" if s.url and s.text != s.url else s.text for s in inline(text)
    )


def _is_web(url: str) -> bool:
    return url.startswith(("http://", "https://", "mailto:"))


def _watch_url(video_id: str) -> str:
    return f"https://www.youtube.com/watch?v={video_id}"


def _body(md: str) -> list[Block]:
    return parse(writer._strip_title_heading(md))


# --- Exporters ---

class Exporter:
    name = ""
    supports_hubs = False
    root: Path

    def __init__(self, cfg: dict):
        self.cfg = cfg

    def find_existing(self, video_id: str) -> str | None:
        raise NotImplementedError

    def write_note(self, title: str, video_id: str, note_md: str) -> str:
        raise NotImplementedError

    def write_tags(self, tags: list[str]) -> list[str]:
        return []

    def record_failure(self, url: str, reason: str) -> None:
        writer.append_failure(self.root, url, reason)

    def _folder(self, name: str) -> str:
        return config.folder(name, self.cfg)


def _root(cfg: dict) -> Path:
    path = (cfg["output"].get("path") or "").strip()
    if not path:
        raise config.ConfigError(
            f"output.path is not set for target '{cfg['output']['target']}'. "
            "Run `afterwatch settings` and pick a folder.")
    return Path(path).expanduser()


class ObsidianExporter(Exporter):
    name = "obsidian"
    supports_hubs = True

    def __init__(self, cfg: dict):
        super().__init__(cfg)
        self.root = _root(cfg)

    def find_existing(self, video_id: str) -> str | None:
        found = writer.find_existing_note(self.root, video_id)
        return str(found) if found else None

    def write_note(self, title: str, video_id: str, note_md: str) -> str:
        return str(writer.write_note(self.root, title, video_id, note_md))

    def write_tags(self, tags: list[str]) -> list[str]:
        return writer.create_tag_stubs(self.root, tags)


class MarkdownExporter(ObsidianExporter):
    """Same layout and wikilinks as the vault, but with standard image links and
    a plain link instead of the YouTube iframe, so any Markdown tool can open it."""
    name = "markdown"

    def write_note(self, title: str, video_id: str, note_md: str) -> str:
        notes = self.root / self._folder("notes")
        images = self.root / self._folder("images")
        rel = urllib.parse.quote(Path(os.path.relpath(images, notes)).as_posix())
        note_md = _EMBED.sub(lambda m: f"![]({rel}/{urllib.parse.quote(m[1].strip())})", note_md)
        path = Path(super().write_note(title, video_id, note_md))
        text = path.read_text(encoding="utf-8")
        iframe = writer._references_section(video_id)
        if text.endswith(iframe):
            link = f"\n\n# References\n[Watch on YouTube]({_watch_url(video_id)})\n"
            path.write_text(text[: -len(iframe)] + link, encoding="utf-8")
        return str(path)


class _Indexed(Exporter):
    """Exporters whose output the vault scan can't see: a JSON index of
    video id -> location stands in for duplicate detection."""

    def _index_file(self) -> Path:
        return config.state_dir() / f"exported-{self.name}.json"

    def _index(self) -> dict:
        try:
            return json.loads(self._index_file().read_text(encoding="utf-8"))
        except (OSError, ValueError):
            return {}

    def _remember(self, video_id: str, location: str) -> None:
        index = self._index()
        index[video_id] = location
        self._index_file().parent.mkdir(parents=True, exist_ok=True)
        self._index_file().write_text(json.dumps(index, indent=2), encoding="utf-8")

    def find_existing(self, video_id: str) -> str | None:
        return self._index().get(video_id)


class _FileExporter(_Indexed):
    ext = ""

    def __init__(self, cfg: dict):
        super().__init__(cfg)
        self.root = _root(cfg)

    def _save(self, path: Path, title: str, video_id: str, note_md: str) -> None:
        raise NotImplementedError

    def find_existing(self, video_id: str) -> str | None:
        found = super().find_existing(video_id)
        return found if found and Path(found).exists() else None

    def write_note(self, title: str, video_id: str, note_md: str) -> str:
        notes = self.root / self._folder("notes")
        notes.mkdir(parents=True, exist_ok=True)
        base = writer.sanitize_filename(title)
        path = notes / f"{base}{self.ext}"
        if path.exists():  # same title, different video
            path = notes / f"{base} ({video_id}){self.ext}"
        self._save(path, title, video_id, note_md)
        self._remember(video_id, str(path))
        return str(path)


class TxtExporter(_FileExporter):
    name = "txt"
    ext = ".txt"

    def _save(self, path: Path, title: str, video_id: str, note_md: str) -> None:
        path.write_text(render_txt(title, note_md, video_id), encoding="utf-8")


def render_txt(title: str, note_md: str, video_id: str) -> str:
    lines = [title, "=" * len(title)]
    prev_tight = False
    for b in _body(note_md):
        tight = b.kind in ("bullet", "number")
        if not (tight and prev_tight):
            lines.append("")
        prev_tight = tight
        pad = "  " * b.level
        if b.kind == "heading":
            text = _plain(b.text)
            lines.append(text)
            if b.level <= 2:
                lines.append(("=" if b.level == 1 else "-") * len(text))
        elif b.kind == "bullet":
            lines.append(f"{pad}- {_plain(b.text)}")
        elif b.kind == "number":
            lines.append(f"{pad}{b.num}. {_plain(b.text)}")
        elif b.kind == "code":
            lines += ["    " + ln for ln in b.text.split("\n")]
        elif b.kind == "quote":
            lines += ["> " + _plain(ln) for ln in b.text.split("\n")]
        elif b.kind == "table":
            rows = [[_plain(c) for c in r] for r in b.rows]
            width = max(len(r) for r in rows)
            rows = [r + [""] * (width - len(r)) for r in rows]
            w = [max(len(r[c]) for r in rows) for c in range(width)]
            fmt = [" | ".join(cell.ljust(w[c]) for c, cell in enumerate(r)).rstrip() for r in rows]
            lines += [fmt[0], "-+-".join("-" * n for n in w)] + fmt[1:]
        elif b.kind == "image":
            lines.append(f"[Screenshot: {b.text}]")
        elif b.kind == "hr":
            lines.append("-" * 40)
        else:
            lines.append(_plain(b.text))
    lines += ["", f"Watch on YouTube: {_watch_url(video_id)}"]
    return "\n".join(lines) + "\n"


def _docx():
    try:
        import docx
    except ImportError as exc:
        raise config.ConfigError(
            "Word (.docx) output needs python-docx: pip install afterwatch[docx]") from exc
    return docx


class DocxExporter(_FileExporter):
    name = "docx"
    ext = ".docx"

    def __init__(self, cfg: dict):
        super().__init__(cfg)
        _docx()

    def _save(self, path: Path, title: str, video_id: str, note_md: str) -> None:
        build_docx(title, note_md, video_id, self.root / self._folder("images")).save(str(path))


def _add_runs(paragraph, text: str, bold: bool = False) -> None:
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.opc.constants import RELATIONSHIP_TYPE

    for s in inline(text):
        if s.url and _is_web(s.url):
            rid = paragraph.part.relate_to(s.url, RELATIONSHIP_TYPE.HYPERLINK, is_external=True)
            link = OxmlElement("w:hyperlink")
            link.set(qn("r:id"), rid)
            run, props, t = OxmlElement("w:r"), OxmlElement("w:rPr"), OxmlElement("w:t")
            for tag, val in (("w:b", bold or s.bold), ("w:i", s.italic)):
                if val:
                    props.append(OxmlElement(tag))
            color, under = OxmlElement("w:color"), OxmlElement("w:u")
            color.set(qn("w:val"), "0563C1")
            under.set(qn("w:val"), "single")
            props.append(color)
            props.append(under)
            t.text = s.text
            t.set(qn("xml:space"), "preserve")
            run.append(props)
            run.append(t)
            link.append(run)
            paragraph._p.append(link)
            continue
        label = f"{s.text} ({s.url})" if s.url and s.text != s.url else s.text
        run = paragraph.add_run(label)
        run.bold = bold or s.bold or None
        run.italic = s.italic or None
        if s.code:
            run.font.name = "Consolas"


def build_docx(title: str, note_md: str, video_id: str, images_dir: Path):
    from docx.shared import Inches, Pt

    doc = _docx().Document()
    doc.add_heading(title, 0)
    for b in _body(note_md):
        if b.kind == "heading":
            doc.add_heading(_plain(b.text), min(b.level, 4))
        elif b.kind == "bullet":
            style = "List Bullet" if b.level == 0 else f"List Bullet {min(b.level + 1, 3)}"
            _add_runs(doc.add_paragraph(style=style), b.text)
        elif b.kind == "number":
            p = doc.add_paragraph()
            p.paragraph_format.left_indent = Inches(0.25 * (b.level + 1))
            p.paragraph_format.first_line_indent = Inches(-0.25)
            p.add_run(f"{b.num}. ")
            _add_runs(p, b.text)
        elif b.kind == "code":
            run = doc.add_paragraph().add_run(b.text)
            run.font.name = "Consolas"
            run.font.size = Pt(9)
        elif b.kind == "quote":
            _add_runs(doc.add_paragraph(style="Quote"), b.text)
        elif b.kind == "table":
            width = max(len(r) for r in b.rows)
            table = doc.add_table(rows=len(b.rows), cols=width)
            table.style = "Table Grid"
            for r, row in enumerate(b.rows):
                for c, cell in enumerate(row):
                    _add_runs(table.cell(r, c).paragraphs[0], cell, bold=(r == 0))
        elif b.kind == "image":
            try:
                doc.add_picture(str(images_dir / Path(b.text).name), width=Inches(6))
            except Exception:  # noqa: BLE001 - missing or unreadable file: keep the note
                doc.add_paragraph().add_run(f"[Screenshot not found: {b.text}]").italic = True
        elif b.kind == "hr":
            doc.add_paragraph("* * *")
        else:
            _add_runs(doc.add_paragraph(), b.text)
    url = _watch_url(video_id)
    p = doc.add_paragraph()
    p.add_run("Watch on YouTube: ")
    _add_runs(p, f"[{url}]({url})")
    return doc


# --- Notion ---

def _notion_id(value: str) -> str:
    """A page id from a 32-hex id (dashes optional) or a notion.so page URL."""
    s = re.split(r"[?#]", value.strip())[0].rstrip("/").replace("-", "")
    m = re.search(r"[0-9a-fA-F]{32}$", s)
    if not m:
        raise config.ConfigError(
            "output.notion_parent must be a Notion page URL or its 32-character id "
            f"(got '{value}').")
    h = m[0].lower()
    return f"{h[:8]}-{h[8:12]}-{h[12:16]}-{h[16:20]}-{h[20:]}"


_NOTION_LANGS = {
    "bash", "c", "c++", "c#", "css", "diff", "docker", "go", "html", "java", "javascript",
    "json", "kotlin", "markdown", "php", "powershell", "python", "ruby", "rust", "shell",
    "sql", "swift", "typescript", "xml", "yaml",
}
_LANG_ALIASES = {"js": "javascript", "ts": "typescript", "py": "python", "sh": "shell",
                 "yml": "yaml", "cpp": "c++", "cs": "c#", "md": "markdown"}


def _rich(segs: list[Seg]) -> list[dict]:
    """Notion rich text; items longer than its 2000-character limit are split."""
    out = []
    for s in segs:
        for i in range(0, len(s.text), _NOTION_TEXT_LIMIT):
            text = {"content": s.text[i:i + _NOTION_TEXT_LIMIT]}
            if s.url and _is_web(s.url) and len(s.url) <= 2000:
                text["link"] = {"url": s.url}
            item: dict = {"type": "text", "text": text}
            if s.bold or s.italic or s.code:
                item["annotations"] = {"bold": s.bold, "italic": s.italic, "code": s.code}
            out.append(item)
    return out


def _block(kind: str, **fields) -> dict:
    return {"object": "block", "type": kind, kind: fields}


def notion_blocks(blocks: list[Block], upload=lambda name: None) -> list[dict]:
    """Top-level Notion blocks. List nesting is kept to one level because a single
    request allows only two; deeper items become children of the top-level item.
    upload(name) returns a file-upload id for a screenshot, or None if it's missing."""
    out: list[dict] = []
    parent: dict | None = None
    for b in blocks:
        if b.kind in ("bullet", "number"):
            kind = "bulleted_list_item" if b.kind == "bullet" else "numbered_list_item"
            item = _block(kind, rich_text=_rich(inline(b.text)))
            if b.level and parent is not None:
                parent[parent["type"]].setdefault("children", []).append(item)
            else:
                out.append(item)
                parent = item
            continue
        parent = None
        if b.kind == "heading":
            out.append(_block(f"heading_{min(b.level, 3)}", rich_text=_rich(inline(b.text))))
        elif b.kind == "code":
            lang = _LANG_ALIASES.get(b.lang, b.lang)
            out.append(_block("code", rich_text=_rich([Seg(b.text)]),
                              language=lang if lang in _NOTION_LANGS else "plain text"))
        elif b.kind == "quote":
            out.append(_block("quote", rich_text=_rich(inline(b.text))))
        elif b.kind == "hr":
            out.append(_block("divider"))
        elif b.kind == "table":
            width = max(len(r) for r in b.rows)
            rows = [_block("table_row", cells=[_rich(inline(c)) for c in r] + [[]] * (width - len(r)))
                    for r in b.rows]
            out.append(_block("table", table_width=width, has_column_header=True, children=rows))
        elif b.kind == "image":
            fid = upload(b.text)
            if fid:
                out.append(_block("image", type="file_upload", file_upload={"id": fid}))
            else:
                out.append(_block("paragraph", rich_text=_rich(
                    [Seg(f"[Screenshot not found: {b.text}]", italic=True)])))
        else:
            out.append(_block("paragraph", rich_text=_rich(inline(b.text))))
    return out


def _retry_after(exc: urllib.error.HTTPError) -> float:
    try:
        return min(float(exc.headers.get("Retry-After") or 1), 60)
    except (ValueError, AttributeError):
        return 1


def _notion_error(exc: urllib.error.HTTPError) -> Exception:
    try:
        message = json.loads(exc.read()).get("message") or exc.reason
    except (ValueError, OSError, AttributeError):
        message = exc.reason
    if exc.code in (401, 403):
        return config.ConfigError(
            f"Notion refused the request ({exc.code}): check the token, and that the parent "
            f"page is shared with your integration. Notion says: {message}")
    return ExportError(f"Notion error {exc.code}: {message}")


class NotionExporter(_Indexed):
    name = "notion"

    def __init__(self, cfg: dict):
        super().__init__(cfg)
        self.root = config.state_dir() / "notion-images"
        self.parent = _notion_id(cfg["output"].get("notion_parent") or "")
        self.token = os.environ.get("NOTION_TOKEN") or config.api_key("notion", cfg)
        if not self.token:
            raise config.ConfigError(
                "no Notion token. Set the NOTION_TOKEN environment variable "
                "(an internal integration secret from notion.so/profile/integrations).")

    def record_failure(self, url: str, reason: str) -> None:
        writer.append_failure(config.state_dir(), url, reason)

    def _api(self, method: str, path: str, body: dict | None = None,
             raw: bytes | None = None, content_type: str = "application/json") -> dict:
        data = raw if raw is not None else (json.dumps(body).encode() if body is not None else None)
        for attempt in range(4):
            req = urllib.request.Request(NOTION_API + path, data=data, method=method, headers={
                "Authorization": f"Bearer {self.token}",
                "Notion-Version": NOTION_VERSION,
                "Content-Type": content_type,
            })
            try:
                with urllib.request.urlopen(req, timeout=60) as resp:
                    return json.loads(resp.read())
            except urllib.error.HTTPError as exc:
                if exc.code in (429, 529) and attempt < 3:
                    time.sleep(_retry_after(exc))
                    continue
                raise _notion_error(exc) from None
            except urllib.error.URLError as exc:
                raise ExportError(f"could not reach Notion: {exc.reason}") from exc
        raise AssertionError("unreachable")

    def _upload(self, name: str) -> str | None:
        """Send one screenshot through Notion's File Upload API; its id goes in an image block."""
        path = self.root / self._folder("images") / Path(name).name
        if not path.is_file():
            return None
        ctype = mimetypes.guess_type(path.name)[0] or "image/jpeg"
        created = self._api("POST", "/file_uploads", {"filename": path.name, "content_type": ctype})
        boundary = uuid.uuid4().hex
        head = (f'--{boundary}\r\nContent-Disposition: form-data; name="file"; '
                f'filename="{path.name}"\r\nContent-Type: {ctype}\r\n\r\n').encode()
        self._api("POST", f"/file_uploads/{created['id']}/send",
                  raw=head + path.read_bytes() + f"\r\n--{boundary}--\r\n".encode(),
                  content_type=f"multipart/form-data; boundary={boundary}")
        return created["id"]

    def write_note(self, title: str, video_id: str, note_md: str) -> str:
        blocks = notion_blocks(_body(note_md), self._upload)
        url = _watch_url(video_id)
        blocks.append(_block("paragraph", rich_text=_rich([Seg("Watch on YouTube", url=url)])))
        page = self._api("POST", "/pages", {
            "parent": {"page_id": self.parent},
            "properties": {"title": {"title": _rich([Seg(title[:_NOTION_TEXT_LIMIT])])}},
            "children": blocks[:_NOTION_BATCH],
        })
        for i in range(_NOTION_BATCH, len(blocks), _NOTION_BATCH):
            self._api("PATCH", f"/blocks/{page['id']}/children",
                      {"children": blocks[i:i + _NOTION_BATCH]})
        self._remember(video_id, page["url"])
        return page["url"]


_TARGETS = {"obsidian": ObsidianExporter, "markdown": MarkdownExporter, "docx": DocxExporter,
            "txt": TxtExporter, "notion": NotionExporter}


def get(cfg: dict | None = None) -> Exporter:
    cfg = cfg or config.load()
    target = cfg["output"]["target"]
    if target not in _TARGETS:
        raise config.ConfigError(
            f"unknown output target '{target}' (choose one of: {', '.join(_TARGETS)})")
    return _TARGETS[target](cfg)
